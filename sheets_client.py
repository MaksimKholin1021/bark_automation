"""Local Excel workbook for Bark leads.

Name and time stay in the cells. The long lead text is an Excel note
on the name cell (Insert > Note), not a column of its own.
"""

from __future__ import annotations

import html
import json
import logging
import re
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Iterable
from zoneinfo import ZoneInfo

import gspread
from google.oauth2.service_account import Credentials
from gspread.utils import rowcol_to_a1
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment

from bark_scraper import Lead

logger = logging.getLogger(__name__)

HEADERS = ["Name", "Added At"]
NOTE_AUTHOR = "Bark"
GOOGLE_SCOPES = ["https://www.googleapis.com/auth/spreadsheets"]
CENTRAL = ZoneInfo("America/Chicago")
# Daily numbering starts at the next Central midnight. 0 through 40 covers 41 accounts.
DAILY_NO_START = datetime(2026, 10, 21, 0, 0, tzinfo=CENTRAL)
ACCOUNT_COUNT = 41
NUMBER_STATE_PATH = Path(__file__).resolve().parent / "lead_numbers.json"


class LocalSheet:
    def __init__(self, path: str | Path):
        self.path = Path(path)
        self.ensure_workbook()

    def ensure_workbook(self) -> None:
        if not self.path.exists():
            self.path.parent.mkdir(parents=True, exist_ok=True)
            wb = Workbook()
            ws = wb.active
            ws.title = "Leads"
            ws.append(HEADERS)
            ws.column_dimensions["A"].width = 28
            ws.column_dimensions["B"].width = 22
            wb.save(self.path)
            logger.info("Created local sheet %s", self.path)
            return
        self._move_content_into_notes()

    def _move_content_into_notes(self) -> None:
        """Move a Content column into notes on the Name cell, then drop that column."""
        wb = load_workbook(self.path)
        ws = wb.active
        header = [str(cell.value or "").strip() for cell in ws[1]]
        lowered = [value.lower() for value in header]
        if "content" not in lowered:
            wb.close()
            return

        content_col = lowered.index("content") + 1
        name_col = lowered.index("name") + 1 if "name" in lowered else 1
        moved = 0
        for row in range(2, ws.max_row + 1):
            name_cell = ws.cell(row, name_col)
            text = str(ws.cell(row, content_col).value or "").strip()
            if not text:
                continue
            if name_cell.comment is None or not str(name_cell.comment.text or "").strip():
                name_cell.comment = _note(text)
                moved += 1
        ws.delete_cols(content_col)
        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 22
        wb.save(self.path)
        wb.close()
        _repair_notes(self.path)
        logger.info("Moved %s content cell(s) into Excel notes", moved)

    def existing_keys(self) -> set[str]:
        wb = load_workbook(self.path)
        try:
            ws = wb.active
            keys: set[str] = set()
            for i, row in enumerate(ws.iter_rows(min_col=1, max_col=1)):
                cell = row[0]
                name = str(cell.value or "").strip()
                if i == 0 and name.lower() == "name":
                    continue
                content = ""
                if cell.comment is not None:
                    content = str(cell.comment.text or "").strip()
                if not name and not content:
                    continue
                keys.add(f"{name.lower()}|{content.lower()}")
            return keys
        finally:
            wb.close()

    def append_leads(self, leads: Iterable[Lead]) -> int:
        rows = list(leads)
        if not rows:
            return 0
        wb = load_workbook(self.path)
        ws = wb.active
        stamp = datetime.now().strftime("%Y-%m-%d %H:%M")
        for lead in rows:
            ws.append([sheet_name(lead), stamp])
            ws.cell(ws.max_row, 1).comment = _note(lead.content)
        wb.save(self.path)
        wb.close()
        _repair_notes(self.path)
        logger.info("Appended %s new lead(s) to %s with notes", len(rows), self.path)
        return len(rows)


def _repair_notes(path: Path) -> None:
    """Put the lead text inside each Excel note popup.

    openpyxl stores the text beside the note and leaves the popup blank,
    so Excel and other viewers show a note with no content.
    """
    path = Path(path)
    with zipfile.ZipFile(path, "r") as zin:
        files = {name: zin.read(name) for name in zin.namelist()}

    comment_name = next(name for name in files if name.startswith("xl/comments/comment"))
    vml_name = next(name for name in files if "commentsDrawing" in name)
    comment_xml = files[comment_name].decode("utf-8")
    texts = [
        html.unescape(text)
        for text in re.findall(r"<t(?:\s[^>]*)?>(.*?)</t>", comment_xml, flags=re.DOTALL)
    ]

    next_id = 1026

    def _shape_id(_match: re.Match[str]) -> str:
        nonlocal next_id
        value = f'shapeId="{next_id}"'
        next_id += 1
        return value

    comment_xml = re.sub(r'shapeId="\d+"', _shape_id, comment_xml)
    files[comment_name] = comment_xml.encode("utf-8")

    vml = files[vml_name].decode("utf-8")
    pattern = re.compile(r'<div style="text-align:left"(?:\s*/>|>.*?</div>)', re.DOTALL)
    matches = list(pattern.finditer(vml))
    if len(matches) == len(texts):
        pieces: list[str] = []
        last = 0
        for match, text in zip(matches, texts):
            pieces.append(vml[last:match.start()])
            pieces.append(f'<div style="text-align:left">{html.escape(text)}</div>')
            last = match.end()
        pieces.append(vml[last:])
        files[vml_name] = "".join(pieces).encode("utf-8")

    tmp = path.with_suffix(".notes.tmp.xlsx")
    with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for name, data in files.items():
            zout.writestr(name, data)
    tmp.replace(path)


class GoogleSheet:
    """Append Bark leads to a Google Sheet.

    Name goes in the Name column. The full lead text is a note on the
    Description cell (Insert > Note), not the cell value.
    """

    def __init__(self, credentials_path: str | Path, spreadsheet_id: str, gid: int = 0):
        self.credentials_path = Path(credentials_path)
        self.spreadsheet_id = spreadsheet_id
        self.gid = gid
        if not self.credentials_path.exists():
            raise FileNotFoundError(f"Missing Google credentials at {self.credentials_path}")
        creds = Credentials.from_service_account_file(str(self.credentials_path), scopes=GOOGLE_SCOPES)
        self._client = gspread.authorize(creds)
        try:
            spreadsheet = self._client.open_by_key(spreadsheet_id)
        except (gspread.exceptions.APIError, PermissionError) as exc:
            email = json.loads(self.credentials_path.read_text(encoding="utf-8")).get("client_email", "")
            cause = exc.__cause__ or exc
            raise PermissionError(
                f"{cause} Share the sheet with {email} as Editor."
            ) from exc
        try:
            self.worksheet = spreadsheet.get_worksheet_by_id(gid)
        except Exception:
            self.worksheet = spreadsheet.sheet1
        (
            self.name_col,
            self.description_col,
            self.created_col,
            self.updated_col,
            self.phone_col,
            self.email_col,
        ) = self._locate_columns()
        self._rows: list[dict] | None = None
        logger.info(
            "Connected to Google Sheet %s tab %s",
            spreadsheet.title,
            self.worksheet.title,
        )

    def _locate_columns(self) -> tuple[int, int, int, int, int, int]:
        headers = [cell.strip() for cell in self.worksheet.row_values(1)]
        lowered = [cell.lower() for cell in headers]
        required = ("name", "description", "created", "updated", "phone", "email")
        missing = [name for name in required if name not in lowered]
        if missing:
            raise RuntimeError(
                "The Google Sheet is missing columns: "
                + ", ".join(missing)
                + f". Found: {headers or '(empty row)'}"
            )
        self.tag_col = lowered.index("no") + 1 if "no" in lowered else 0
        return tuple(lowered.index(name) + 1 for name in required)

    def begin_cycle(self) -> None:
        self._rows = None

    def is_already_listed(self, lead: Lead) -> bool:
        """True when this Bark lead is already stored, even if the name was edited."""
        return bool(self.matching_rows(lead))

    def matching_rows(self, lead: Lead) -> list[dict]:
        """Sheet rows for this Bark lead.

        "Siddharth Ray" is the same lead as Bark's "Siddharth" when the city matches.
        An empty starred placeholder such as "Richard *" is not.
        A starred row that still has this lead's note, such as "St *", is.
        """
        self._snapshot()
        phone, email = _phone_and_email(lead.content)
        place = (lead.location or "").split(",")[0].strip().lower()
        matches: list[dict] = []
        for row in self._rows or []:
            if row.get("upwork"):
                continue
            note = (row.get("note") or "").strip()
            starred = "*" in row["name"]
            if starred and not note:
                continue
            note_name = note.splitlines()[0].strip() if note else ""
            if not _names_match(row["name"], lead.name) and not _names_match(note_name, lead.name):
                continue
            blob = f"{row['name']}\n{note}".lower()
            same_contact = bool(
                phone and email and row["phone"] == phone and row["email"] == email
            )
            if starred:
                if (place and place in blob) or same_contact:
                    matches.append(row)
                continue
            if _compact_name(row["name"]) == _compact_name(lead.name):
                matches.append(row)
                continue
            if (place and place in blob) or same_contact:
                matches.append(row)
        return matches

    def is_recent_listing(self, row: dict) -> bool:
        """True when this row is part of the latest leads, not an older namesake."""
        self._snapshot()
        if not self._rows:
            return False
        last = self._rows[-1]["index"]
        return row["index"] >= last - 40

    def _snapshot(self) -> None:
        if self._rows is not None:
            return
        names = self.worksheet.col_values(self.name_col)
        phones = self.worksheet.col_values(self.phone_col)
        emails = self.worksheet.col_values(self.email_col)
        tags = self.worksheet.col_values(self.tag_col) if self.tag_col else []
        notes = self.worksheet.get_notes()
        desc_index = self.description_col - 1
        rows: list[dict] = []
        for index, name in enumerate(names):
            if index == 0:
                continue
            note = ""
            if index < len(notes) and desc_index < len(notes[index]):
                note = str(notes[index][desc_index] or "")
            phone_cell = phones[index] if index < len(phones) else ""
            email_cell = emails[index] if index < len(emails) else ""
            phone, email = _phone_and_email("\n".join([str(phone_cell), str(email_cell), note]))
            tag = tags[index].strip().lower() if index < len(tags) else ""
            rows.append(
                {
                    "index": index,
                    "name": name.strip(),
                    "phone": phone,
                    "email": email,
                    "note": note,
                    "upwork": tag == "u",
                }
            )
        self._rows = rows

    def names_needing_notes(self) -> set[str]:
        """First names of rows whose description note is not the full lead panel yet."""
        names = self.worksheet.col_values(self.name_col)
        notes = self.worksheet.get_notes()
        tags = self.worksheet.col_values(self.tag_col) if self.tag_col else []
        desc_index = self.description_col - 1
        wanted: set[str] = set()
        for index, name in enumerate(names):
            if index == 0 or not name.strip():
                continue
            if index < len(tags) and tags[index].strip().lower() == "u":
                continue
            current = ""
            if index < len(notes) and desc_index < len(notes[index]):
                current = str(notes[index][desc_index] or "").strip()
            if "Details" in current and len(current) > 200:
                continue
            parts = _name_parts(name)
            if parts:
                wanted.add(parts[0])
        return wanted

    def existing_keys(self) -> set[str]:
        names = self.worksheet.col_values(self.name_col)
        notes = self.worksheet.get_notes()
        keys: set[str] = set()
        desc_index = self.description_col - 1
        for i, name in enumerate(names):
            name = name.strip()
            if i == 0 and name.lower() == "name":
                continue
            content = ""
            if i < len(notes) and desc_index < len(notes[i]):
                content = str(notes[i][desc_index] or "").strip()
            if not name and not content:
                continue
            keys.add(f"{name.lower()}|{content.lower()}")
        return keys

    def append_leads(self, leads: Iterable[Lead]) -> tuple[int, list[tuple[int, int | str]]]:
        rows = list(leads)
        if not rows:
            return 0, []
        name_used = len(self.worksheet.col_values(self.name_col))
        desc_used = len(self.worksheet.col_values(self.description_col))
        start = max(name_used, desc_used, 1) + 1
        end = start + len(rows) - 1
        name_range = f"{rowcol_to_a1(start, self.name_col)}:{rowcol_to_a1(end, self.name_col)}"
        self.worksheet.update([[sheet_name(lead)] for lead in rows], range_name=name_range)
        numbers: list[int | str] = [""] * len(rows)
        if self.tag_col:
            numbered = [index for index, lead in enumerate(rows) if not _is_exist_lead(lead)]
            if numbered:
                assigned, day_state = self._next_lead_numbers(len(numbered), start)
                for index, number in zip(numbered, assigned):
                    numbers[index] = number
                if day_state is not None:
                    _save_number_state(*day_state)
            number_range = f"{rowcol_to_a1(start, self.tag_col)}:{rowcol_to_a1(end, self.tag_col)}"
            self.worksheet.update([[number] for number in numbers], range_name=number_range)
        notes = {
            rowcol_to_a1(start + offset, self.description_col): lead.content
            for offset, lead in enumerate(rows)
        }
        self.worksheet.insert_notes(notes)
        placed = [(start + offset, number) for offset, number in enumerate(numbers)]
        logger.info("Appended %s new lead(s)", len(rows))
        return len(rows), placed

    def name_rows(self) -> list[dict]:
        """Name, No, and city for every lead row. City comes from the description note."""
        names = self.worksheet.col_values(self.name_col)
        notes = self.worksheet.get_notes()
        tags = self.worksheet.col_values(self.tag_col) if self.tag_col else []
        desc_index = self.description_col - 1
        rows: list[dict] = []
        for index, name in enumerate(names):
            if index == 0:
                continue
            note = ""
            if index < len(notes) and desc_index < len(notes[index]):
                note = str(notes[index][desc_index] or "")
            number = tags[index].strip() if index < len(tags) else ""
            rows.append(
                {
                    "row": index + 1,
                    "name": str(name or "").strip(),
                    "no": number,
                    "city": _city_from_note(note),
                }
            )
        return rows

    def full_name_is_other_city(self, row: int, full_name: str, rows: list[dict] | None = None) -> bool:
        """True when this exact full name is already stored for a different city.

        A previous Carl Cannon in Detroit is not the full name for a new Carl
        in Port Allen. The Bark email for the new lead has to arrive on its own.
        """
        key = _name_key(full_name)
        if not key or len(_plain_name(full_name).split()) < 2:
            return False
        indexed = rows if rows is not None else self.name_rows()
        target = next((item for item in indexed if item["row"] == row), None)
        target_city = str(target["city"] if target else "").lower()
        for item in indexed:
            if item["row"] == row or _name_key(item["name"]) != key:
                continue
            other_city = str(item["city"] or "").lower()
            if not target_city or not other_city or other_city != target_city:
                return True
        return False

    def write_full_name(
        self,
        row: int,
        short_name: str,
        full_name: str,
        rows: list[dict] | None = None,
    ) -> bool:
        """Replace a contacted lead's first name with the full name from Bark's email."""
        current = str(self.worksheet.cell(row, self.name_col).value or "").strip()
        if not _same_short_name(current, short_name):
            logger.warning(
                "Row %s is %r, so the full name %r was not written over %r",
                row,
                current,
                full_name,
                short_name,
            )
            return False
        full = re.sub(r"\s+", " ", (full_name or "").strip())
        if _name_key(full) == _name_key(current):
            if len(_plain_name(current).split()) != 1 or "*" in current:
                return True
            updated = _mark_same_name(current)
        else:
            if self.full_name_is_other_city(row, full, rows):
                logger.info(
                    "Left row %s as %r. %r already belongs to a lead in another city.",
                    row,
                    current,
                    full,
                )
                return False
            updated = _merge_full_name(current, full)
        if updated == current:
            return True
        self.worksheet.update_acell(rowcol_to_a1(row, self.name_col), updated)
        logger.info("Updated row %s name %r -> %r", row, current, updated)
        return True

    def _next_lead_numbers(self, count: int, start_row: int) -> tuple[list[int | str], tuple[str, int] | None]:
        """Number new leads for the current Central day.

        From 2026-10-21 00:00 Central, the first lead of each day is 0 and the
        41st is 40. Later leads that day keep an empty No cell. Before that
        midnight, numbering still continues through 40 and then wraps to 0.
        """
        now = datetime.now(CENTRAL)
        if now >= DAILY_NO_START:
            day = now.date().isoformat()
            assigned = _load_number_state(day)
            numbers, assigned = _daily_lead_numbers(assigned, count)
            return numbers, (day, assigned)

        tags = self.worksheet.col_values(self.tag_col)
        last = -1
        for index, value in enumerate(tags[1:], start=2):
            if index >= start_row:
                break
            text = str(value).strip()
            if text.isdigit() and 0 <= int(text) <= 40:
                last = int(text)
        numbers: list[int | str] = []
        for _ in range(count):
            last = 0 if last >= 40 else last + 1
            numbers.append(last)
        return numbers, None

    def update_detail_notes(self, leads: Iterable[Lead]) -> int:
        """Replace short description notes with the full right-hand lead panel."""
        names = self.worksheet.col_values(self.name_col)
        notes = self.worksheet.get_notes()
        tags = self.worksheet.col_values(self.tag_col) if self.tag_col else []
        desc_index = self.description_col - 1
        by_row: dict[int, list[Lead]] = {}
        for lead in leads:
            if "\n" not in (lead.content or ""):
                continue
            matches = [
                index
                for index, name in enumerate(names)
                if index > 0
                and "*" not in name
                and not (index < len(tags) and tags[index].strip().lower() == "u")
                and _name_parts(name) == _name_parts(lead.name)
            ]
            if len(matches) != 1:
                continue
            by_row.setdefault(matches[0], []).append(lead)
        updates: dict[str, str] = {}
        for index, group in by_row.items():
            if len(group) != 1:
                continue
            lead = group[0]
            current = ""
            if index < len(notes) and desc_index < len(notes[index]):
                current = str(notes[index][desc_index] or "").strip()
            if not _should_replace_note(current, lead, single=True):
                continue
            updates[rowcol_to_a1(index + 1, self.description_col)] = lead.content
        if not updates:
            return 0
        self.worksheet.insert_notes(updates)
        logger.info("Updated %s description note(s) with the full lead panel", len(updates))
        return len(updates)

def _daily_lead_numbers(assigned: int, count: int) -> tuple[list[int | str], int]:
    """Give out 0..40, then blanks. `assigned` is how many numbers today already used."""
    numbers: list[int | str] = []
    for _ in range(count):
        if assigned >= ACCOUNT_COUNT:
            numbers.append("")
        else:
            numbers.append(assigned)
            assigned += 1
    return numbers, assigned


def _load_number_state(day: str) -> int:
    if not NUMBER_STATE_PATH.exists():
        return 0
    try:
        data = json.loads(NUMBER_STATE_PATH.read_text(encoding="utf-8"))
    except Exception as exc:
        logger.warning("Could not read %s: %s", NUMBER_STATE_PATH, exc)
        return 0
    if data.get("day") != day:
        return 0
    assigned = data.get("assigned", 0)
    if isinstance(assigned, int) and assigned >= 0:
        return assigned
    return 0


def _save_number_state(day: str, assigned: int) -> None:
    NUMBER_STATE_PATH.write_text(
        json.dumps({"day": day, "assigned": assigned}, indent=2),
        encoding="utf-8",
    )


def _name_parts(value: str) -> list[str]:
    value = value.lower().replace("*", " ")
    value = re.sub(r"\(.*?\)", " ", value)
    value = re.sub(r"\barchitect\b", " ", value)
    skip = {"str", "arch"}
    return [part for part in re.findall(r"[a-z0-9']+", value) if part not in skip]


def _phone_and_email(text: str) -> tuple[str, str]:
    phone = ""
    email = ""
    for line in (text or "").splitlines():
        compact = re.sub(r"\s+", "", line).lower()
        if not email and "@" in compact:
            email = compact
        if not phone and (
            re.search(r"\(\d{3}\)", line) or re.search(r"\d{3}[-.\s]*\d{3}[-.\s]*\d{4}", line)
        ):
            phone = compact
    return phone, email


def _compact_name(value: str) -> str:
    """Ignore spaces and punctuation so 'Gary docter' matches 'Garydocter'."""
    value = re.sub(r"\(str\)", "", value or "", flags=re.I)
    return re.sub(r"[^a-z0-9]", "", value.lower())


def _is_exist_lead(lead: Lead) -> bool:
    return "(exist)" in (lead.name or "").lower()


def _same_short_name(current: str, short_name: str) -> bool:
    return _first_token(current) == _first_token(short_name) and bool(_first_token(short_name))


def _first_token(value: str) -> str:
    cleaned = re.sub(r"\(.*?\)", " ", value or "")
    cleaned = cleaned.replace("*", " ")
    parts = re.findall(r"[a-z0-9']+", cleaned.lower())
    return parts[0] if parts else ""


def _plain_name(value: str) -> str:
    text = re.sub(r"\(.*?\)", " ", value or "")
    text = text.replace("*", " ")
    return re.sub(r"\s+", " ", text).strip()


def _name_key(value: str) -> str:
    return _plain_name(value).lower()


def _city_from_note(note: str) -> str:
    for line in (note or "").splitlines():
        match = re.match(r"^(.+?),\s*([A-Z]{2})\b", line.strip())
        if not match:
            continue
        city = match.group(1).strip()
        if len(city) < 2 or any(char.isdigit() for char in city):
            continue
        return f"{city}, {match.group(2)}"
    return ""


def _mark_same_name(current: str) -> str:
    """Bark's email used the same name already in the cell, so mark it."""
    suffix = ""
    match = re.search(r"\s*(\((?:str|exist)\))\s*$", current or "", flags=re.I)
    base = current or ""
    if match:
        suffix = " " + match.group(1)
        base = base[: match.start()]
    base = _plain_name(base)
    return f"{base} (*){suffix}".strip()


def _merge_full_name(current: str, full_name: str) -> str:
    full = re.sub(r"\s+", " ", (full_name or "").strip())
    if "(str)" in (current or "").lower() and "(str)" not in full.lower():
        return f"{full} (str)"
    return full


def exist_display_name(name: str) -> str:
    cleaned = re.sub(r"\s*\(exist\)\s*$", "", (name or "").strip(), flags=re.I).strip()
    return f"{cleaned} (exist)"


def sheet_name(lead: Lead) -> str:
    """Structural leads are marked '(str)'. Architectural leads keep the name only."""
    name = (lead.name or "").strip()
    if _is_exist_lead(lead):
        return name
    if not _is_structural(lead) or "(str)" in name.lower():
        return name
    return f"{name} (str)"


def _is_structural(lead: Lead) -> bool:
    category = (lead.category or "").strip().lower()
    if "structural" in category:
        return True
    if "architectural" in category:
        return False
    for line in (lead.content or "").splitlines():
        lower = line.strip().lower()
        if lower.startswith("structural engineer"):
            return True
        if lower.startswith("architectural services"):
            return False
    return False


def _names_match(sheet_name: str, bark_name: str) -> bool:
    """Bark often shows a first name; the sheet may already have the full name."""
    sheet_compact = _compact_name(sheet_name)
    bark_compact = _compact_name(bark_name)
    if sheet_compact and bark_compact and sheet_compact == bark_compact:
        return True
    sheet_parts = _name_parts(sheet_name)
    bark_parts = _name_parts(bark_name)
    if not sheet_parts or not bark_parts:
        return False
    if sheet_parts == bark_parts:
        return True
    return len(bark_parts) == 1 and sheet_parts[0] == bark_parts[0]


def _should_replace_note(current: str, lead: Lead, *, single: bool) -> bool:
    if "Details" in current and len(current) > 200:
        return False
    if not single:
        place = lead.location.split(",")[0].strip().lower()
        if place and place not in current.lower():
            return False
    return True


def _note(text: str) -> Comment:
    clean = (text or "").replace("\x00", "").strip()
    # Excel notes are the popup from Insert > Note.
    comment = Comment(clean, NOTE_AUTHOR)
    comment.width = 320
    comment.height = 160
    return comment
